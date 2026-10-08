"""Read-only workbook ingestion with retained source rows and explicit knowledge gaps.

No identical-looking print is deleted. Vendor trade IDs are absent, so equality
of cell values does not establish that two exchange executions are duplicates.
"""
from __future__ import annotations

import csv
import gzip
import hashlib
import json
import math
import re
import zipfile
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo
from xml.etree import ElementTree as ET

from openpyxl import load_workbook

FLOW_HEADERS = ["Date", "Symbol", "DTE", "Option", "Quantity", "Price", "Side", "Premium", "Sentiment", "Open Interest", "QTY > OI ?", "Expired?"]
ANGLE_HEADERS = ["Date", "Ticker", "Company Name", "Hidden Angle", "Why It Matters"]
SUMMARY_HEADERS = ["#", "Ticker", "Bull Premium", "Bear Premium", "New Position $", "Bull/Bear", "# Bull Trades", "Hidden Angle", "Why It Matters"]
_NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
_MONTHS = {m.lower(): i for i, m in enumerate("Jan Feb Mar Apr May Jun Jul Aug Sep Oct Nov Dec".split(), 1)}


def json_value(value: Any) -> Any:
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        return None
    return value


def sha256(path: str | Path) -> str:
    h = hashlib.sha256()
    with Path(path).open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def parse_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if not isinstance(value, str) or not value.strip():
        return None
    text = value.strip()
    try:
        return datetime.fromisoformat(text.replace("Z", "+00:00")).date()
    except ValueError:
        pass
    for fmt in ("%m/%d/%Y", "%d-%b-%Y", "%d %B %Y", "%B %d, %Y", "%d.%m.%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    return None


def number(value: Any) -> float | None:
    if isinstance(value, bool) or value is None:
        return None
    try:
        result = float(str(value).replace(",", "").replace("$", "").strip())
        return result if math.isfinite(result) else None
    except (ValueError, TypeError):
        return None


def ticker_tokens(value: Any) -> list[str]:
    """Tokenize the source's ticker field, never infer tickers from company text."""
    text = str(value or "").upper().strip()
    tokens = re.split(r"\s*(?:/|,|;|\||\s+&\s+)\s*", text)
    return list(dict.fromkeys(t.strip().lstrip("$") for t in tokens if re.fullmatch(r"\$?[A-Z^][A-Z0-9.^-]{0,14}", t.strip())))


def infer_direction(side: Any, option_type: str | None) -> tuple[str | None, str]:
    label = str(side or "").strip().lower()
    if option_type not in ("CALL", "PUT"):
        return None, "option type unavailable"
    if label in ("near ask", "on ask", "above ask", "at ask", "ask"):
        return ("BULLISH" if option_type == "CALL" else "BEARISH"), "inferred ask-side activity, opening status unconfirmed"
    if label in ("near bid", "on bid", "below bid", "at bid", "bid"):
        return ("BEARISH" if option_type == "CALL" else "BULLISH"), "inferred bid-side activity, opening status unconfirmed"
    return None, "trade side ambiguous or unavailable"


def expiry_fields(trade_date: date | None, dte: float | None, option: Any) -> dict:
    """Use reported calendar DTE and cross-check contract text; do not invent a root map."""
    text = str(option or "")
    m = re.search(r"\b(Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)(\d{2})(?:\s+(\d{1,2})(?:st|nd|rd|th)\b)?", text, re.I)
    expiry = trade_date + timedelta(days=int(dte)) if trade_date and dte is not None and dte.is_integer() and dte >= 0 else None
    explicit = None
    if m and m.group(3):
        try:
            explicit = date(2000 + int(m.group(2)), _MONTHS[m.group(1).lower()], int(m.group(3)))
        except ValueError:
            pass
    flags = []
    if expiry and m:
        if expiry.year != 2000 + int(m.group(2)) or expiry.month != _MONTHS[m.group(1).lower()]:
            flags.append("DTE-derived expiry conflicts with contract month/year")
        if explicit and explicit != expiry:
            flags.append("DTE-derived expiry conflicts with explicit contract day")
    if not expiry:
        flags.append("expiry unavailable from trade date and DTE")
    return {"expiry": expiry.isoformat() if expiry else None,
            "expiry_basis": "trade_date_plus_vendor_DTE" if expiry else None,
            "explicit_text_expiry": explicit.isoformat() if explicit else None,
            "contract_root": text.split()[0] if text.strip() else None,
            "expiry_validation_flags": flags,
            "expiry_independently_verified": False}


def _last_nonempty_row(path: Path, ws: Any) -> int:
    """Bound styled million-row sheets by actual XML values, including preserved formulas."""
    last = 0
    with zipfile.ZipFile(path) as archive:
        with archive.open(ws._worksheet_path.lstrip("/")) as f:
            for _, el in ET.iterparse(f, events=("end",)):
                if el.tag == _NS + "row":
                    if any(any(n.text for n in c.iter() if n.tag in (_NS + "v", _NS + "t")) for c in el.findall(_NS + "c")):
                        last = int(el.attrib["r"])
                    el.clear()
    return last


def _raw_id(name: str, digest: str, sheet: str, row: int, columns: int) -> str:
    return f"{name}:{digest[:12]}:{sheet}!A{row}:{chr(64 + columns)}{row}"


def _fingerprint(values: list[Any]) -> str:
    return hashlib.sha256(json.dumps(values, ensure_ascii=False, sort_keys=True, default=json_value).encode("utf-8")).hexdigest()


def _date_range(rows: list[dict], key: str) -> dict:
    dates = [r[key] for r in rows if r.get(key)]
    return {"first": min(dates, default=None), "last": max(dates, default=None)}


def ingest_sources(flow_path: str | Path, research_path: str | Path, observed_at: str | None = None) -> dict:
    """Parse actual files. observed_at proves this ingest only, not historical availability."""
    flow_path, research_path = Path(flow_path), Path(research_path)
    observed_at = observed_at or datetime.now(timezone.utc).isoformat()
    if datetime.fromisoformat(observed_at.replace("Z", "+00:00")).tzinfo is None:
        raise ValueError("observed_at must include a timezone")
    flow_records, research_records, summary_records, manifests = [], [], [], []
    sheets_audit = []
    for kind, path in (("flow", flow_path), ("research", research_path)):
        digest = sha256(path)
        wb = load_workbook(path, read_only=True, data_only=True)
        manifest = {"file": path.name, "sha256": digest, "byte_count": path.stat().st_size,
                    "observed_ingestion_at": observed_at, "historical_available_at": None,
                    "historical_availability_proven": False, "sheets": []}
        for ws in wb:
            is_flow = ws.title in ("Flow Analytics (Stocks)", "Flow Analytics (ETFs)")
            is_angle = ws.title == "Hidden Angles Database"
            is_summary = ws.title == "Hidden Angle Flow Analysis"
            last = _last_nonempty_row(path, ws)
            metadata_rows = [list(row) for row in ws.iter_rows(min_row=1, max_row=min(last, 5), max_col=min(ws.max_column, 12), values_only=True)]
            meta = {"sheet": ws.title, "declared_max_row": ws.max_row, "last_nonempty_xml_row": last,
                    "max_columns": ws.max_column, "metadata_rows": [[json_value(v) for v in row] for row in metadata_rows]}
            manifest["sheets"].append(meta)
            if not (is_flow or is_angle or is_summary):
                meta["status"] = "Unrecognized sheet retained in raw workbook; no signal extracted"
                continue
            headers, header_row, max_col = (FLOW_HEADERS, 5, 12) if is_flow else ((ANGLE_HEADERS, 4, 5) if is_angle else (SUMMARY_HEADERS, 4, 9))
            actual_header = list(next(ws.iter_rows(min_row=header_row, max_row=header_row, max_col=max_col, values_only=True)))
            if actual_header != headers:
                raise ValueError(f"Unexpected schema in {path.name}/{ws.title}: {actual_header!r}")
            records = []
            fingerprints = Counter()
            for excel_row, row in enumerate(ws.iter_rows(min_row=header_row + 1, max_row=last, max_col=max_col, values_only=True), header_row + 1):
                if not any(value is not None for value in row):
                    continue
                raw = [json_value(v) for v in row]
                source_id = _raw_id(path.name, digest, ws.title, excel_row, max_col)
                fp = _fingerprint(raw)
                fingerprints[fp] += 1
                record = {"source_id": source_id, "source_file": path.name, "source_sha256": digest,
                          "source_sheet": ws.title, "source_row": excel_row,
                          "raw_values": dict(zip(headers, raw)), "raw_row_fingerprint": fp,
                          "source_available_at": None, "availability_status": "historical time unknown",
                          "observed_ingestion_at": observed_at}
                if is_flow:
                    td, dte = parse_date(row[0]), number(row[2])
                    option_type = "CALL" if re.search(r"\bCalls?\b", str(row[3]), re.I) else ("PUT" if re.search(r"\bPuts?\b", str(row[3]), re.I) else None)
                    inferred, inference = infer_direction(row[6], option_type)
                    supplied = str(row[8] or "").upper().strip()
                    supplied = supplied if supplied in ("BULLISH", "BEARISH") else None
                    direction = inferred if inferred and (not supplied or supplied == inferred) else None
                    q, p, premium, oi = number(row[4]), number(row[5]), number(row[7]), number(row[9])
                    expected = q * p * 100 if q is not None and p is not None else None
                    error = premium - expected if premium is not None and expected is not None else None
                    flags = []
                    if not td: flags.append("missing or invalid trade date")
                    if premium is None or premium < 0: flags.append("missing or invalid premium")
                    if q is None or q <= 0: flags.append("missing or invalid contract quantity")
                    if p is None or p <= 0: flags.append("missing or invalid option price")
                    if not ticker_tokens(row[1]): flags.append("missing or invalid exact source ticker")
                    if supplied and inferred and supplied != inferred: flags.append("supplied sentiment conflicts with trade-side/type inference")
                    if error is not None and abs(error) > max(1.0, abs(expected) * 1e-5): flags.append("premium differs from quantity x price x assumed100 multiplier")
                    record.update({"asset_category": "stock" if ws.title.endswith("(Stocks)") else "ETF_or_index_context",
                                   "trade_date": td.isoformat() if td else None, "ticker": str(row[1] or "").strip().upper().lstrip("$"),
                                   "option": str(row[3] or ""), "option_type": option_type, "reported_trade_dte": dte,
                                   "quantity": q, "option_price_usd": p, "premium_usd": premium,
                                   "assumed_multiplier": 100, "expected_premium_usd": expected, "premium_arithmetic_error_usd": error,
                                   "side": row[6], "supplied_sentiment": supplied, "inferred_sentiment": inferred,
                                   "direction": direction, "direction_basis": inference, "open_interest": oi,
                                   "quantity_exceeds_oi": q > oi if q is not None and oi is not None else None,
                                   "opening_trade_confirmed": False, "validation_flags": flags})
                    record.update(expiry_fields(td, dte, row[3]))
                    if record["contract_root"] and record["contract_root"] != record["ticker"].lstrip("^"):
                        record["validation_flags"].append("contract root differs from stock/index field; exact mapping unverified")
                    flow_records.append(record)
                elif is_angle:
                    rd = parse_date(row[0])
                    record.update({"research_date": rd.isoformat() if rd else None, "date_basis": "vendor row date, publication time unverified",
                                   "tickers": ticker_tokens(row[1]), "company_name": row[2], "hidden_angle": row[3], "why_it_matters": row[4],
                                   "original_publisher": None, "original_source_url": None, "event_date_verified": None,
                                   "direction_assessment": "requires evidence review", "holding_window_assessment": "not yet assessed",
                                   "thesis_verified": False})
                    research_records.append(record)
                else:
                    record.update({"tickers": ticker_tokens(row[1]), "bull_premium_usd": number(row[2]), "bear_premium_usd": number(row[3]),
                                   "vendor_new_position_usd": number(row[4]), "vendor_bull_bear_ratio": number(row[5]),
                                   "vendor_bull_trade_count": number(row[6]), "hidden_angle": row[7], "why_it_matters": row[8],
                                   "aggregation_scope": "cumulative unexpired forward-position sample, not Monday daily flow",
                                   "research_date": None, "thesis_verified": False})
                    summary_records.append(record)
                records.append(record)
            duplicate_groups = sum(1 for n in fingerprints.values() if n > 1)
            duplicate_looking_excess = sum(n - 1 for n in fingerprints.values() if n > 1)
            for r in records:
                r["identical_looking_group_size"] = fingerprints[r["raw_row_fingerprint"]]
            sheets_audit.append({"file": path.name, "sheet": ws.title, "record_count": len(records),
                                 "duplicate_looking_groups": duplicate_groups, "duplicate_looking_excess_rows_retained": duplicate_looking_excess,
                                 "removed_rows": 0, "date_range": _date_range(records, "trade_date" if is_flow else "research_date")})
        wb.close()
        manifests.append(manifest)
    etf_context_tickers = {r["ticker"] for r in flow_records if r["asset_category"] == "ETF_or_index_context"}
    for record in flow_records:
        record["asset_classification_basis"] = "vendor worksheet, exact broker settlement/instrument class unverified"
        record["underlying_share_identity_verified"] = False
        record["also_appears_in_ETF_index_sheet"] = record["ticker"] in etf_context_tickers
    sessions = defaultdict(lambda: {"rows": 0, "bull_premium_usd": 0.0, "bear_premium_usd": 0.0,
                                     "unclassified_premium_usd": 0.0, "tickers": set(), "expiry_counts": Counter()})
    flow_as_of = max((r["trade_date"] for r in flow_records if r["trade_date"]), default=None)
    expiry_asof_counts = Counter()
    for r in flow_records:
        if r["trade_date"]:
            s = sessions[(r["asset_category"], r["trade_date"])]
            s["rows"] += 1; s["tickers"].add(r["ticker"]); s["expiry_counts"][r["expiry"] or "unknown"] += 1
            s[{"BULLISH": "bull_premium_usd", "BEARISH": "bear_premium_usd"}.get(r["direction"], "unclassified_premium_usd")] += r["premium_usd"] or 0
        if r["expiry"] and flow_as_of:
            delta = (date.fromisoformat(r["expiry"]) - date.fromisoformat(flow_as_of)).days
            expiry_asof_counts["expired_before_snapshot" if delta < 0 else ("expires_on_snapshot" if delta == 0 else "unexpired_after_snapshot")] += 1
        else:
            expiry_asof_counts["unknown_expiry"] += 1
    session_rows = []
    for (cat, session), s in sorted(sessions.items()):
        session_rows.append({"category": cat, "session": session, "rows": s["rows"], "ticker_count": len(s["tickers"]),
                             "bull_premium_usd": s["bull_premium_usd"], "bear_premium_usd": s["bear_premium_usd"],
                             "net_bull_premium_usd": s["bull_premium_usd"] - s["bear_premium_usd"],
                             "unclassified_premium_usd": s["unclassified_premium_usd"], "expiry_counts": dict(s["expiry_counts"])})
    flag_counts = Counter(flag for r in flow_records for flag in r["validation_flags"] + r["expiry_validation_flags"])
    stock_root_flags = [r for r in flow_records if r["asset_category"] == "stock" and any("contract root" in f for f in r["validation_flags"])]
    historical_rows = [r for r in flow_records if r["trade_date"] and flow_as_of and r["trade_date"] < flow_as_of]
    has_expired = expiry_asof_counts.get("expired_before_snapshot", 0) > 0
    audit = {"audit_version": "Rev12_source_audit_1", "observed_ingestion_at": observed_at,
             "source_manifest": manifests, "sheets": sheets_audit, "flow_records": len(flow_records),
             "research_records": len(research_records), "summary_records": len(summary_records),
             "research_rows_containing_replacement_characters": sum("\ufffd" in str(r.get("hidden_angle", "")) + str(r.get("why_it_matters", "")) for r in research_records),
             "flow_date_range": _date_range(flow_records, "trade_date"), "research_date_range": _date_range(research_records, "research_date"),
             "session_totals": session_rows, "expiry_status_at_latest_flow_session": dict(expiry_asof_counts),
             "validation_flag_counts": dict(flag_counts), "historical_rows_before_latest_session": len(historical_rows),
             "flow_missing_normalized_fields": {key: sum(r.get(key) is None or r.get(key) == "" for r in flow_records) for key in ("trade_date", "ticker", "reported_trade_dte", "expiry", "quantity", "option_price_usd", "premium_usd", "direction", "open_interest")},
             "research_missing_normalized_fields": {key: sum(not r.get(key) for r in research_records) for key in ("research_date", "tickers", "company_name", "hidden_angle", "why_it_matters")},
             "flow_side_labels": dict(Counter(str(r.get("side")) for r in flow_records)),
             "flow_premium_below_declared_minimum_rows": sum(r.get("premium_usd") is not None and r["premium_usd"] < 100_000 for r in flow_records),
             "stock_sheet_contract_root_mapping_flag_rows": len(stock_root_flags),
             "stock_sheet_tickers_also_in_ETF_index_sheet": sorted({r["ticker"] for r in flow_records if r["asset_category"] == "stock" and r["also_appears_in_ETF_index_sheet"]}),
             "instrument_classification_limit": "The vendor Stocks worksheet also contains fund-like symbols. Worksheet membership does not verify underlying-stock identity. Cross-sheet ETF/index matches are excluded from the weekly stock filter; every other ticker still needs broker identity and settlement checks.",
             "survivorship_assessment": ("Some expired records remain; completeness still unproven" if has_expired else "No expired contract found in retained historical rows. Consistent with survivor-filtering; complete historical short-DTE backtest unavailable."),
             "source_availability_assessment": "Dates and as-of captions are not receipt timestamps. Monday/Tues point-in-time reconstruction requires archived files with proven receipt times.",
             "duplicate_policy": "All raw rows retained. Identical cell values do not prove duplicate executions. No vendor transaction IDs supplied.",
             "premium_units": "USD option premium; arithmetic checked against quantity x option price x assumed100 shares/contract. Adjusted-contract multiplier remains unverified.",
             "research_source_assessment": "Vendor-compiled thesis text has row dates but no original URLs/publication timestamps. Summary totals are cumulative forward-position history, not the Monday session."}
    return {"flow_records": flow_records, "research_records": research_records, "summary_records": summary_records, "audit": audit}


def write_audit(bundle: dict, out_dir: str | Path) -> list[str]:
    """Save machine-readable source rows and useful audit controls without modifying inputs."""
    out = Path(out_dir); out.mkdir(parents=True, exist_ok=True)
    paths = []
    for key in ("flow_records", "research_records", "summary_records"):
        path = out / f"Normalized_{key}.jsonl.gz"
        # gzip mtime=0 makes normalized snapshots repeatable; timestamps remain explicit in records.
        with path.open("wb") as raw:
            with gzip.GzipFile(filename="", mode="wb", fileobj=raw, mtime=0) as gz:
                for record in bundle[key]:
                    gz.write((json.dumps(record, ensure_ascii=False, default=json_value) + "\n").encode("utf-8"))
        paths.append(str(path))
    path = out / "Source_Audit.json"
    path.write_text(json.dumps(bundle["audit"], indent=2, ensure_ascii=False, default=json_value), encoding="utf-8"); paths.append(str(path))
    path = out / "Flow_Session_Totals.csv"
    fields = ["category", "session", "rows", "ticker_count", "bull_premium_usd", "bear_premium_usd", "net_bull_premium_usd", "unclassified_premium_usd", "expiry_counts"]
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fields); writer.writeheader()
        for row in bundle["audit"]["session_totals"]:
            writer.writerow({**row, "expiry_counts": json.dumps(row["expiry_counts"], sort_keys=True)})
    paths.append(str(path))
    a = bundle["audit"]
    local_ingestion = datetime.fromisoformat(a["observed_ingestion_at"].replace("Z", "+00:00")).astimezone(ZoneInfo("Europe/Zurich")).isoformat()
    lines = ["# Revision 12 source audit", "", f"Observed ingestion: {local_ingestion} (Swiss local time). Historical receipt times remain unknown.", "",
             "| Source sheet | Records | First date | Last date | Identical-looking excess retained |", "|---|---:|---|---|---:|"]
    for s in a["sheets"]:
        lines.append(f"| {s['file']} / {s['sheet']} | {s['record_count']:,} | {s['date_range']['first'] or 'not supplied'} | {s['date_range']['last'] or 'not supplied'} | {s['duplicate_looking_excess_rows_retained']:,} |")
    lines.extend(["", a["duplicate_policy"], "", a["survivorship_assessment"], "", f"Expiry controls as of {a['flow_date_range']['last']}: `{json.dumps(a['expiry_status_at_latest_flow_session'], sort_keys=True)}`.", "", a["research_source_assessment"], "", a["premium_units"], "", a["source_availability_assessment"], "",
                  "The full daily controls are in Flow_Session_Totals.csv. All normalized rows carry file hash, sheet, Excel row, source ID and the original cell values. No row deletion, market prices, simulated profit or original-source verification is implied."])
    path = out / "Source_Audit.md"; path.write_text("\n".join(lines) + "\n", encoding="utf-8"); paths.append(str(path))
    return paths
